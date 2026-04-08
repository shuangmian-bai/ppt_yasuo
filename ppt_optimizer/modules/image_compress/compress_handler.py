"""图片压缩模块示例"""

import os
import io
from enum import Enum
from pptx import Presentation
from PIL import Image
from concurrent.futures import ThreadPoolExecutor, as_completed
from tqdm import tqdm


# 处理结果状态枚举
class ProcessStatus(Enum):
    SUCCESS = "正常处理完成"
    ERROR_INTERRUPTED = "异常中断"
    FILE_NOT_FOUND = "文件不存在"


def compress_images_in_pptx(input_pptx, output_pptx, max_pixel=1920, quality=85, max_workers=4):
    """
    压缩PPTX中的所有图片
    
    Args:
        input_pptx: 输入PPTX文件路径
        output_pptx: 输出PPTX文件路径
        max_pixel: 最大像素尺寸
        quality: JPEG压缩质量
        max_workers: 线程数量
    
    Returns:
        tuple: (ProcessStatus状态, 压缩数量或错误信息)
    """
    if not os.path.exists(input_pptx):
        return ProcessStatus.FILE_NOT_FOUND, f"输入文件不存在: {input_pptx}"
    
    output_dir = os.path.dirname(output_pptx)
    if output_dir:
        try:
            os.makedirs(output_dir, exist_ok=True)
        except Exception as e:
            return ProcessStatus.ERROR_INTERRUPTED, f"创建输出目录失败: {str(e)}"
    
    try:
        prs = Presentation(input_pptx)
        total_compressed = 0
        
        tasks = [(slide, idx) for idx, slide in enumerate(prs.slides, 1)]
        total_slides = len(tasks)
        
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_slide = {
                executor.submit(process_slide_compress, slide, idx, max_pixel, quality): idx
                for slide, idx in tasks
            }
            
            with tqdm(total=total_slides, desc="压缩进度", unit="页") as pbar:
                for future in as_completed(future_to_slide):
                    slide_idx, compressed = future.result()
                    total_compressed += compressed
                    pbar.set_postfix({"已压缩": total_compressed})
                    pbar.update(1)
        
        prs.save(output_pptx)
        
        return ProcessStatus.SUCCESS, total_compressed
        
    except Exception as e:
        return ProcessStatus.ERROR_INTERRUPTED, f"处理过程中发生错误: {str(e)}"


def process_slide_compress(slide, slide_idx, max_pixel=1920, quality=85):
    """处理单个幻灯片中的图片压缩"""
    compressed = 0
    shapes = list(slide.shapes)
    
    for shape in reversed(shapes):
        if shape.shape_type == 13:  # 图片
            img = shape.image
            ext = img.ext.lower()
            
            # 只处理位图格式
            if ext in ["jpg", "jpeg", "png", "bmp", "gif"]:
                try:
                    img_bytes = img.blob
                    img_pil = Image.open(io.BytesIO(img_bytes)).convert("RGB")
                    
                    w, h = img_pil.size
                    if max(w, h) > max_pixel:
                        scale = max_pixel / max(w, h)
                        new_w = int(w * scale)
                        new_h = int(h * scale)
                        img_pil = img_pil.resize((new_w, new_h), Image.Resampling.LANCZOS)
                    
                    output_io = io.BytesIO()
                    img_pil.save(output_io, format="JPEG", optimize=True, quality=quality)
                    output_io.seek(0)
                    
                    left = shape.left
                    top = shape.top
                    width = shape.width
                    height = shape.height
                    
                    slide.shapes._spTree.remove(shape._element)
                    
                    slide.shapes.add_picture(
                        output_io,
                        left, top, width, height
                    )
                    compressed += 1
                    
                except Exception as e:
                    continue
    
    return slide_idx, compressed
