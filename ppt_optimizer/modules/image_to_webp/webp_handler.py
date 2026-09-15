import os
import io
import threading
from enum import Enum

from pptx import Presentation
from pptx.parts.image import Image as PptxImage
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
import pptx.opc.spec as spec
from PIL import Image
from concurrent.futures import ThreadPoolExecutor, as_completed
from tqdm import tqdm


# ---------------------------------------------------------------------------
# 给 python-pptx 打补丁：使其支持写入 WebP
# python-pptx 的 Image.ext 只认 BMP/GIF/JPEG/PNG/TIFF/WMF，会导致 add_picture(webp) 报错。
# 只需在 ext_map 里加入 WEBP，并在 image_content_types 里注册 "webp": "image/webp" 即可。
# ---------------------------------------------------------------------------
if "webp" not in spec.image_content_types:
    spec.image_content_types["webp"] = "image/webp"

    def _ext(self):
        ext_map = {
            "BMP": "bmp",
            "GIF": "gif",
            "JPEG": "jpg",
            "PNG": "png",
            "TIFF": "tiff",
            "WMF": "wmf",
            "WEBP": "webp",
        }
        fmt = self._format
        if fmt not in ext_map:
            raise ValueError("unsupported image format, got '%s'" % fmt)
        return ext_map[fmt]

    PptxImage.ext = property(_ext)


# 处理结果状态枚举
class ProcessStatus(Enum):
    SUCCESS = "正常处理完成"
    ERROR_INTERRUPTED = "异常中断"
    FILE_NOT_FOUND = "文件不存在"


def _has_alpha(img):
    """判断 PIL 图片是否含透明通道"""
    return img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)


def _encode_webp(img, quality):
    """将 PIL 图片编码为 WebP 字节；含 alpha 的走无损，其余走有损压缩"""
    output_io = io.BytesIO()
    if _has_alpha(img):
        img = img.convert("RGBA")
        img.save(output_io, format="WEBP", lossless=True)
    else:
        img = img.convert("RGB")
        img.save(output_io, format="WEBP", quality=quality, method=6)
    output_io.seek(0)
    return output_io


def process_slide_to_webp(slide, slide_idx, max_pixel=1920, quality=80, lock=None):
    """将单个幻灯片中的全部图片转换为 WebP 并等比压缩。

    直接在 <a:blip> 层面替换 r:embed 指向，而非「删旧图 + add_picture」，
    因此能统一覆盖三类图片，且不破坏形状结构（裁剪、效果、旋转、分组、图片填充等）：
      1. 顶层图片 <p:pic>
      2. 分组内的图片 <p:grpSp> 中的 <p:pic>
      3. 自动形状/文本框里的图片填充 <p:sp> 的 <a:blipFill>
    """
    converted = 0

    for blip in list(slide._element.iter(qn("a:blip"))):
        rId = blip.get(qn("r:embed"))
        if rId is None:
            continue  # 外链图片（r:link）或无内嵌，跳过

        try:
            # 直接从 OPC 关系取原始字节，兼容 webp/mpo 等 python-pptx 读不了的格式
            blob = slide.part.related_part(rId).blob
            img_pil = Image.open(io.BytesIO(blob))

            # 等比缩放（只缩小不放大）
            w, h = img_pil.size
            if max(w, h) > max_pixel:
                scale = max_pixel / max(w, h)
                img_pil = img_pil.resize(
                    (int(w * scale), int(h * scale)), Image.Resampling.LANCZOS
                )

            # 编码为 WebP
            output_io = _encode_webp(img_pil, quality)

            # get_or_add_image_part 会改包级图片注册表 + 关系图，需加锁避免多线程竞态。
            # 返回 (image_part, rId)；相同字节会复用已有 part/rId，天然去重。
            if lock is not None:
                lock.acquire()
            try:
                _, new_rId = slide.part.get_or_add_image_part(output_io)
            finally:
                if lock is not None:
                    lock.release()

            # 把该 blip 的引用切到新的 webp part（形状结构、位置、裁剪完全保留）
            blip.set(qn("r:embed"), new_rId)
            converted += 1

        except Exception:
            continue

    return slide_idx, converted


def _drop_orphan_image_rels(slide):
    """删除幻灯片中不再被任何元素引用的图片关系，使孤立图片 part 在保存时被丢弃。

    收集所有 r:embed 引用（含 <a:blip> 与 SVG 的 <asvg:svgBlip>），
    只删除 reltype 为 image 且不再被引用的关系，避免误删 SVG/其它关系。
    """
    used = set()
    for el in slide._element.iter():
        rId = el.get(qn("r:embed"))
        if rId:
            used.add(rId)

    rels = slide.part.rels
    for rId in list(rels.keys()):
        if rels[rId].reltype == RT.IMAGE and rId not in used:
            rels.pop(rId)


def convert_images_to_webp_in_pptx(input_pptx, output_pptx, max_pixel=1920, quality=80, max_workers=2):
    """
    将 PPTX 中全部图片转为 WebP 并等比压缩

    Args:
        input_pptx: 输入 PPTX 文件路径
        output_pptx: 输出 PPTX 文件路径
        max_pixel: 长边最大像素尺寸
        quality: WebP 有损压缩质量 (1-100)
        max_workers: 线程数量

    Returns:
        tuple: (ProcessStatus状态, 转换数量或错误信息)
    """
    if not os.path.exists(input_pptx):
        return ProcessStatus.FILE_NOT_FOUND, f"输入文件不存在: {input_pptx}"

    output_dir = os.path.dirname(output_pptx)
    if output_dir:
        try:
            os.makedirs(output_dir, exist_ok=True)
        except Exception as e:
            return ProcessStatus.ERROR_INTERRUPTED, f"创建输出目录失败: {str(e)}"

    try:
        prs = Presentation(input_pptx)
        total_converted = 0
        lock = threading.Lock()

        tasks = [(slide, idx) for idx, slide in enumerate(prs.slides, 1)]
        total_slides = len(tasks)

        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_slide = {
                executor.submit(process_slide_to_webp, slide, idx, max_pixel, quality, lock): idx
                for slide, idx in tasks
            }

            with tqdm(total=total_slides, desc="WebP转换进度", unit="页") as pbar:
                for future in as_completed(future_to_slide):
                    slide_idx, converted = future.result()
                    total_converted += converted
                    pbar.set_postfix({"已转换": total_converted})
                    pbar.update(1)

        # 清理孤立图片关系（单线程，避免并发改关系图遍历竞态）
        for slide in prs.slides:
            _drop_orphan_image_rels(slide)

        prs.save(output_pptx)

        return ProcessStatus.SUCCESS, total_converted

    except Exception as e:
        return ProcessStatus.ERROR_INTERRUPTED, f"处理过程中发生错误: {str(e)}"
